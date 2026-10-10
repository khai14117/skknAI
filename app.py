import io
import streamlit as st
import google.generativeai as genai
from docx import Document

# --- CẤU HÌNH TRANG WEB ---
st.set_page_config(
    page_title="Trợ Lý AI Viết Sáng Kiến Kinh Nghiệm (SKKN)",
    page_icon="💡",
    layout="wide",
)

st.title("💡 Trợ Lý AI Hỗ Trợ Viết Sáng Kiến Kinh Nghiệm (SKKN)")
st.caption("Bài tập 4: Tự động hóa xây dựng cấu trúc và nội dung SKKN chuẩn khoa học giáo dục")

# --- LẤY API KEY ẨN TỪ SECRETS CỦA STREAMLIT CLOUD ---
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
except Exception:
    GOOGLE_API_KEY = None

# --- GIAO DIỆN NHẬP LIỆU ---
with st.form("skkn_form"):
    col1, col2 = st.columns(2)
    with col1:
        subject = st.selectbox(
            "📖 Môn học / Lĩnh vực",
            ["Toán học", "Ngữ văn", "Tiếng Anh", "Khoa học tự nhiên", "Vật lý", "Hóa học", "Sinh học", "Lịch sử & Địa lí", "Tin học", "Công nghệ", "Công tác Quản lý / Chủ nhiệm lớp"]
        )
        grade = st.selectbox(
            "🏫 Khối lớp áp dụng",
            ["Khối THCS (Lớp 6, 7, 8, 9)", "Khối THPT (Lớp 10, 11, 12)", "Toàn cấp học"]
        )
    with col2:
        target_audience = st.text_input("🎯 Đối tượng học sinh / Phạm vi thực tế", placeholder="Ví dụ: Học sinh lớp 10 trường THPT...")
        experience_title = st.text_input("💡 Tên đề tài / Vấn đề cốt lõi của SKKN", placeholder="Ví dụ: Một số biện pháp nâng cao hứng thú giải toán hình học không gian...")

    solution_focus = st.text_area(
        "📝 Ý tưởng hoặc giải pháp cốt lõi của bạn (nếu có):",
        placeholder="Mô tả ngắn gọn sáng kiến của bạn (nếu để trống, AI sẽ tự đề xuất phương pháp tối ưu hoàn toàn)."
    , height=100)

    submit_btn = st.form_submit_button("🚀 Biên soạn Sáng Kiến Kinh Nghiệm", type="primary")

# --- HÀM TẠO FILE WORD (.DOCX) ---
def export_to_word(text_content):
    doc = Document()
    doc.add_heading("SÁNG KIẾN KINH NGHIỆM - NGÀNH GIÁO DỤC", level=0)
    
    for paragraph in text_content.split("\n"):
        p_clean = paragraph.strip()
        if not p_clean:
            continue
        if p_clean.startswith("# "):
            doc.add_heading(p_clean.replace("# ", ""), level=1)
        elif p_clean.startswith("## "):
            doc.add_heading(p_clean.replace("## ", ""), level=2)
        elif p_clean.startswith("### "):
            doc.add_heading(p_clean.replace("### ", ""), level=3)
        elif p_clean.startswith("- ") or p_clean.startswith("* "):
            doc.add_paragraph(p_clean[2:], style='List Bullet')
        else:
            doc.add_paragraph(p_clean)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- XỬ LÝ KHI BẤM NÚT ---
if submit_btn:
    if not GOOGLE_API_KEY:
        st.error("⚠️ Hệ thống chưa cấu hình API Key ở mục Secrets trên Streamlit Cloud!")
    elif not experience_title:
        st.error("⚠️ Vui lòng nhập tên đề tài hoặc vấn đề cốt lõi của SKKN!")
    else:
        try:
            genai.configure(api_key=GOOGLE_API_KEY)
            model = genai.GenerativeModel("gemini-2.5-flash")

            prompt = f"""
            Bạn là một chuyên gia nghiên cứu khoa học sư phạm và viết Sáng kiến kinh nghiệm (SKKN) xuất sắc tại Việt Nam.
            Hãy biên soạn một bản Sáng kiến kinh nghiệm hoàn chỉnh, khoa học và mang tính thực tiễn cao cho:
            - Môn học / Lĩnh vực: {subject}
            - Phạm vi: {grade} ({target_audience})
            - Tên đề tài: "{experience_title}"
            - Ý tưởng cốt lõi bổ sung: {solution_focus}

            YÊU CẦU CẤU TRÚC BẮT BUỘC CỦA SKKN:
            1. PHẦN MỞ ĐẦU:
               - Lý do chọn đề tài (Tính cấp thiết, thực trạng vấn đề).
               - Mục đích và nhiệm vụ nghiên cứu.
               - Đối tượng và phạm vi nghiên cứu.
               - Phương pháp nghiên cứu.
            2. PHẦN NỘI DUNG (GIẢI PHÁP THỰC THI):
               - Cơ sở lý luận và thực trạng vấn đề trước khi áp dụng sáng kiến (Những tồn tại, khó khăn).
               - Các biện pháp/giải pháp mới đã tiến hành để giải quyết vấn đề (Trình bày chi tiết 3-4 biện pháp cụ thể, có minh họa ví dụ sư phạm thực tế).
               - Khảo nghiệm tính hiệu quả (Số liệu hoặc minh chứng đánh giá trước và sau khi áp dụng).
            3. PHẦN KẾT LUẬN VÀ KIẾN NGHỊ:
               - Kết luận, ý nghĩa của sáng kiến đối với công tác giảng dạy.
               - Bài học kinh nghiệm và kiến nghị đề xuất.

            Trình bày chi tiết, chuyên nghiệp, chuẩn văn phong báo cáo khoa học giáo dục bằng tiếng Việt.
            """

            with st.spinner("🔄 AI đang nghiên cứu, hệ thống hóa và viết SKKN chuyên sâu..."):
                response = model.generate_content(prompt)
                skkn_text = response.text

            st.success("✅ Đã biên soạn thành công Sáng kiến kinh nghiệm!")
            st.markdown("---")
            st.markdown(skkn_text)

            # Tạo file Word
            word_file = export_to_word(skkn_text)
            
            st.markdown("---")
            st.download_button(
                label="📥 Tải xuống File Word Sáng Kiến Kinh Nghiệm (.docx)",
                data=word_file,
                file_name=f"SKKN_{subject}_{experience_title[:15]}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary"
            )

        except Exception as e:
            st.error(f"❌ Có lỗi xảy ra khi gọi AI: {str(e)}")